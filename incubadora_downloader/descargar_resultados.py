"""Descarga diaria de resultados de nacimiento desde el sistema de incubación.

Automatiza el flujo manual de http://192.168.9.218:50000/NewIncubadora:
  Zona=Incubación, Indicador=Fecha Nacimiento, Origen=Reproductoras,
  Rango de Fechas (un día), Confirmar -> por cada Sala: Sectores -> marcar
  pabellones -> Pabellones Detalle -> Excel.

Cada Excel descargado ya trae identificada la ubicación (Planta/SubPlanta/
Sector/Desc.Sector); este script solo agrega la columna Fecha y consolida
todo en un único archivo Excel acumulado.

Uso:
    python descargar_resultados.py                     # descarga el día de ayer
    python descargar_resultados.py --desde 2026-09-01 --hasta 2026-09-28
    python descargar_resultados.py --visible            # ver el navegador mientras corre
"""

from __future__ import annotations

import argparse
import datetime as dt
import re
import sys
from pathlib import Path

import openpyxl
from playwright.sync_api import Page, sync_playwright

BASE_URL = "http://192.168.9.218:50000/NewIncubadora/servlet/com.incubadora.resultincub"

# Valores fijos confirmados en el formulario (ver README para cómo obtenerlos
# de nuevo si el sistema cambia).
ORIGEN_REPRODUCTORAS = "2"
# Con Origen=Reproductoras, Sexo solo tiene una opción disponible ("Broiler")
# y queda seleccionada automáticamente; no hace falta tocarla.

SCRIPT_DIR = Path(__file__).resolve().parent
DESCARGAS_DIR = SCRIPT_DIR / "descargas_tmp"
MAESTRO_XLSX = SCRIPT_DIR / "nacimientos_acumulado.xlsx"

# Las filas de sector/pabellón dentro de "Sectores" siempre incluyen "PLANTA N".
PLANTA_RE = re.compile(r"PLANTA\s*\d+")

DEBUG = False


MESES = [
    "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
    "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre",
]


def set_select(page: Page, selector: str, texto_visible: str, value: str, reintentos: int = 3) -> None:
    """Selecciona un <select> tecleando su texto visible y verifica el valor.

    select_option() cambia el valor del combo visualmente (input_value()
    lo confirma) pero el evento que dispara NO queda marcado como
    "confiable" para el JavaScript de esta página (GeneXus), así que el
    servidor nunca se entera del cambio real: el HTML de depuración mostró
    el combo en el valor correcto pero el GXState del servidor seguía en
    el valor anterior. Por eso se simula tecleo real (type-ahead nativo
    del <select>) en vez de select_option().
    """
    locator = page.locator(selector)
    for intento in range(reintentos):
        locator.click()
        page.keyboard.type(texto_visible, delay=120)
        page.keyboard.press("Tab")
        page.wait_for_load_state("networkidle")
        # La página sincroniza los cambios por WebSocket, no por HTTP
        # normal, así que "networkidle" no alcanza a esperar esa
        # sincronización; se agrega una espera generosa adicional.
        page.wait_for_timeout(1500)
        if locator.input_value() == value:
            return
    raise RuntimeError(
        f"No se pudo fijar {selector}={value!r} ({texto_visible!r}) tras "
        f"{reintentos} intentos (quedó en {locator.input_value()!r})"
    )


def set_dia(page: Page, selector: str, dia: int, reintentos: int = 3) -> None:
    """Escribe un campo Día simulando tecleo real y verifica el valor final.

    .fill() no dispara la validación gx.num.valid_integer (que depende de
    eventos de teclado) y el campo terminaba reseteado a "0". Se usa
    press_sequentially() y se verifica/reintenta por si acaso.
    """
    locator = page.locator(selector)
    valor = str(dia)
    for intento in range(reintentos):
        locator.click()
        locator.fill("")
        locator.press_sequentially(valor, delay=120)
        page.keyboard.press("Tab")  # dispara el blur que activa la validación
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(1500)
        if locator.input_value() == valor:
            return
    raise RuntimeError(
        f"No se pudo fijar {selector}={valor!r} tras {reintentos} intentos "
        f"(quedó en {locator.input_value()!r})"
    )


def set_origen_reproductoras(page: Page, reintentos: int = 3) -> None:
    """Fija Origen=Reproductoras esperando la cascada real de Sexo.

    Origen dispara, además del onchange normal, un evento especial
    (EVQORIGEN.CLICK) que recalcula las opciones de Sexo en el
    servidor. Con Año/Mes (que no tienen esta cascada) una espera fija
    bastaba, pero para Origen el GXState del servidor seguía mostrando
    "Abuelas" aunque el combo ya se viera en "Reproductoras". Se espera
    a una señal concreta de que la cascada terminó: que Sexo quede con
    la única opción "Broiler" que corresponde a Reproductoras.
    """
    locator = page.locator("#vQORIGEN")
    for intento in range(reintentos):
        locator.click()
        page.keyboard.type("Reproductoras", delay=120)
        page.keyboard.press("Tab")
        page.wait_for_load_state("networkidle")
        try:
            page.wait_for_function(
                """() => {
                    const sel = document.querySelector('#vPSEXO');
                    return !!sel && sel.options.length === 1
                        && sel.options[0].text.trim() === 'Broiler';
                }""",
                timeout=8000,
            )
        except Exception:
            continue
        page.wait_for_timeout(1500)
        if locator.input_value() == ORIGEN_REPRODUCTORAS:
            return
    raise RuntimeError(
        "No se pudo fijar Origen=Reproductoras: la cascada de Sexo nunca "
        "quedó en 'Broiler' tras varios intentos."
    )


def configurar_filtros(page: Page, fecha: dt.date) -> None:
    page.goto(BASE_URL)
    page.wait_for_load_state("networkidle")
    # El radio "Rango de Fechas" dispara una recarga del formulario al
    # hacer clic; se usa .click() en vez de .check() porque Playwright
    # verifica el estado "checked" apenas termina la recarga y a veces
    # todavía no quedó reflejado, aunque el clic sí se aplicó.
    page.get_by_role("radio", name="Rango de Fechas").click()
    page.wait_for_load_state("networkidle")
    page.wait_for_timeout(1500)

    # Origen primero: cambia qué opciones son válidas en otros campos
    # (como Sexo), así que conviene fijarlo antes que el resto.
    set_origen_reproductoras(page)
    set_select(page, "#vQANO", str(fecha.year), str(fecha.year))
    mes_texto = MESES[fecha.month - 1]
    set_select(page, "#vQMESDESDE", mes_texto, str(fecha.month))
    set_select(page, "#vQMESHASTA", mes_texto, str(fecha.month))
    set_dia(page, "#vDIADESDE", fecha.day)
    set_dia(page, "#vDIAHASTA", fecha.day)

    # Espera final extra antes de Confirmar: todos los cambios anteriores
    # se sincronizan por WebSocket, y conviene darles tiempo de sobra a
    # asentarse antes de enviar la consulta.
    page.wait_for_timeout(1500)
    page.get_by_role("button", name="Confirmar").click()
    page.wait_for_load_state("networkidle")
    # La grilla de resultados también se llena por WebSocket después de
    # que la página ya terminó de cargar, así que se espera a que
    # aparezca al menos una fila antes de darse por vencido (si de
    # verdad no hay datos ese día, este wait simplemente agota su plazo).
    try:
        page.wait_for_function(
            """() => {
                const tbody = document.querySelector('#Grid1ContainerTbl tbody');
                return !!tbody && tbody.children.length > 0;
            }""",
            timeout=8000,
        )
    except Exception:
        pass

    if DEBUG:
        page.screenshot(path=str(SCRIPT_DIR / "debug_ultima_consulta.png"), full_page=True)
        (SCRIPT_DIR / "debug_ultima_consulta.html").write_text(page.content(), encoding="utf-8")


def descargar_dia(page: Page, fecha: dt.date) -> list[Path]:
    configurar_filtros(page, fecha)

    # Buscar el checkbox "dentro" de una fila (por rol o por <tr>) resolvía
    # a varios checkboxes de filas distintas a la vez (la grilla parece
    # tener alguna tabla anidada). En cambio, los checkboxes habilitados
    # (sin contar Totales, que viene disabled) corresponden 1 a 1, en el
    # mismo orden, con las filas de datos reales.
    checkboxes_sala = page.locator('input[type="checkbox"]:not([disabled])')
    n_salas = checkboxes_sala.count()
    if n_salas == 0:
        print(f"  Sin datos para {fecha.isoformat()}")
        return []

    archivos: list[Path] = []
    for i in range(n_salas):
        # .click() en vez de .check(): marcar el checkbox dispara una
        # recarga y .check() falla verificando el estado justo después
        # (mismo comportamiento que el radio "Rango de Fechas").
        page.locator('input[type="checkbox"]:not([disabled])').nth(i).click()
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(500)
        if DEBUG:
            page.screenshot(path=str(SCRIPT_DIR / f"debug_tras_check_sala{i}.png"), full_page=True)
        try:
            page.get_by_role("link", name="Sectores").click(timeout=10000)
        except Exception:
            if DEBUG:
                page.screenshot(
                    path=str(SCRIPT_DIR / f"debug_fallo_sectores_sala{i}.png"), full_page=True
                )
                (SCRIPT_DIR / f"debug_fallo_sectores_sala{i}.html").write_text(
                    page.content(), encoding="utf-8"
                )
            raise
        page.wait_for_load_state("networkidle")
        page.locator("tr").filter(has_text=PLANTA_RE).first.wait_for(
            state="visible", timeout=10000
        )

        # Las filas no cambian de posición, pero sus checkboxes se
        # DESHABILITAN (no desaparecen) una vez descargado su grupo.
        # Como "checkboxes habilitados" se achica por cada grupo ya
        # procesado, hay que restarle ese corrimiento a los índices
        # originales de los grupos siguientes (se calculan todos una
        # sola vez, al principio, cuando todavía están todos habilitados).
        filas_sector = page.locator("tr").filter(has_text=PLANTA_RE)
        n_sectores = filas_sector.count()
        grupos: dict[str, list[int]] = {}
        for j in range(n_sectores):
            texto = filas_sector.nth(j).inner_text()
            m = PLANTA_RE.search(texto)
            clave = m.group(0) if m else "UNICA"
            grupos.setdefault(clave, []).append(j)

        corrimiento = 0
        for planta, indices_originales in grupos.items():
            checkboxes_sector = page.locator('input[type="checkbox"]:not([disabled])')
            for idx in indices_originales:
                checkboxes_sector.nth(idx - corrimiento).click()
                page.wait_for_load_state("networkidle")
                page.wait_for_timeout(500)
            page.get_by_role("link", name="Pabellones Detalle").click()
            page.wait_for_load_state("networkidle")
            # Confirmar que de verdad llegamos a la página con el link
            # "Excel" antes de clickearlo (mismo patrón que las otras
            # transiciones: a veces la recarga no terminó de asentarse).
            page.get_by_role("link", name="Excel").wait_for(state="visible", timeout=10000)
            # El archivo parece generarse en el servidor después de que el
            # link ya es visible; se espera un poco más antes de clickear.
            page.wait_for_timeout(2000)
            if DEBUG:
                href = page.get_by_role("link", name="Excel").get_attribute("href")
                print(f"    (debug) href del link Excel: {href!r}")
            with page.expect_download(timeout=20000) as info:
                page.get_by_role("link", name="Excel").click()
            descarga = info.value
            ruta = DESCARGAS_DIR / f"{fecha.isoformat()}_sala{i}_{planta.replace(' ', '')}.xlsx"
            descarga.save_as(ruta)
            archivos.append(ruta)
            page.get_by_role("link", name="Salir").click()
            page.wait_for_load_state("networkidle")
            # Confirmar que de verdad volvimos a la vista de Sectores
            # antes de seguir (si no, el siguiente grupo de planta se
            # intenta marcar sobre la página equivocada).
            page.locator("tr").filter(has_text=PLANTA_RE).first.wait_for(
                state="visible", timeout=10000
            )
            corrimiento += len(indices_originales)

        page.wait_for_timeout(500)
        page.get_by_role("link", name="Salir").click()
        page.wait_for_load_state("networkidle")
        page.wait_for_timeout(1000)
        # Confirmar que de verdad volvimos a la lista principal de salas
        # antes de procesar la siguiente (si el "Salir" no alcanzaba a
        # completarse, el script seguía marcando checkboxes sobre la
        # vista de Sectores de la sala anterior).
        try:
            page.get_by_role("button", name="Confirmar").wait_for(state="visible", timeout=15000)
        except Exception:
            if DEBUG:
                page.screenshot(
                    path=str(SCRIPT_DIR / f"debug_fallo_salir_sala{i}.png"), full_page=True
                )
            raise

    return archivos


def fechas_ya_procesadas() -> set[str]:
    if not MAESTRO_XLSX.exists():
        return set()
    wb = openpyxl.load_workbook(MAESTRO_XLSX, read_only=True)
    try:
        ws = wb.active
        return {row[0] for row in ws.iter_rows(min_row=2, max_col=1, values_only=True) if row[0]}
    finally:
        wb.close()


def consolidar(archivos_con_fecha: list[tuple[dt.date, Path]]) -> None:
    if not archivos_con_fecha:
        return

    if MAESTRO_XLSX.exists():
        wb = openpyxl.load_workbook(MAESTRO_XLSX)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Nacimientos"
        primero = openpyxl.load_workbook(archivos_con_fecha[0][1])
        encabezado = [c.value for c in next(primero.active.iter_rows(min_row=1, max_row=1))]
        ws.append(["Fecha", *encabezado])
        primero.close()

    for fecha, ruta in archivos_con_fecha:
        wb_origen = openpyxl.load_workbook(ruta)
        hoja = wb_origen.active
        for fila in hoja.iter_rows(min_row=2, values_only=True):
            ws.append([fecha.isoformat(), *fila])
        wb_origen.close()

    wb.save(MAESTRO_XLSX)
    wb.close()

    for _, ruta in archivos_con_fecha:
        ruta.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--desde", type=dt.date.fromisoformat, help="YYYY-MM-DD")
    parser.add_argument("--hasta", type=dt.date.fromisoformat, help="YYYY-MM-DD")
    parser.add_argument(
        "--visible",
        dest="headless",
        action="store_false",
        default=True,
        help="Mostrar el navegador mientras corre (por defecto corre oculto)",
    )
    parser.add_argument(
        "--debug",
        action="store_true",
        help="Guardar debug_ultima_consulta.png/.html con lo que se ve tras Confirmar",
    )
    args = parser.parse_args()

    global DEBUG
    DEBUG = args.debug

    hoy = dt.date.today()
    desde = args.desde or (hoy - dt.timedelta(days=1))
    hasta = args.hasta or desde

    DESCARGAS_DIR.mkdir(exist_ok=True)
    ya_procesadas = fechas_ya_procesadas()

    with sync_playwright() as pw:
        browser = pw.chromium.launch(headless=args.headless)
        context = browser.new_context(accept_downloads=True)
        page = context.new_page()
        # Por si marcar un checkbox u otra acción dispara un diálogo de
        # confirmación del navegador (alert/confirm): sin esto, Playwright
        # se queda esperando indefinidamente a que alguien lo responda.
        page.on("dialog", lambda dialog: dialog.accept())

        fecha = desde
        while fecha <= hasta:
            iso = fecha.isoformat()
            if iso in ya_procesadas:
                print(f"{iso}: ya estaba en el acumulado, se omite.")
                fecha += dt.timedelta(days=1)
                continue

            print(f"Procesando {iso}...")
            try:
                archivos = descargar_dia(page, fecha)
                consolidar([(fecha, a) for a in archivos])
                print(f"  OK: {len(archivos)} archivo(s) consolidado(s).")
            except Exception as exc:  # noqa: BLE001 - se registra y sigue con el resto de los días
                print(f"  ERROR en {iso}: {exc}", file=sys.stderr)

            fecha += dt.timedelta(days=1)

        context.close()
        browser.close()


if __name__ == "__main__":
    main()
